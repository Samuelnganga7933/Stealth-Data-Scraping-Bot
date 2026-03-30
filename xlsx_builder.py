"""
output/xlsx_builder.py — Styled XLSX report builder
Takes scraped data, outputs clean tabular Excel files
Multiple sheets: one per data type / platform
"""

import os
from datetime import datetime
from pathlib import Path
from typing import Optional
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from loguru import logger

OUTPUT_DIR = Path("output/files")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# ── Colours ───────────────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")
ALT_ROW_FILL  = PatternFill("solid", fgColor="EBF3FB")
HEADER_FONT   = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT     = Font(name="Arial", size=10)
TITLE_FONT    = Font(name="Arial", bold=True, size=14, color="1F4E79")
THIN_BORDER   = Border(
    left=Side(style="thin", color="BDD7EE"),
    right=Side(style="thin", color="BDD7EE"),
    top=Side(style="thin", color="BDD7EE"),
    bottom=Side(style="thin", color="BDD7EE"),
)


class XLSXBuilder:

    def __init__(self, output_dir: Path = OUTPUT_DIR):
        self.output_dir = output_dir

    def build(self, data: dict, filename: str = None) -> Path:
        """
        Build an XLSX from a dict of {sheet_name: list_of_dicts}.
        Returns path to saved file.
        """
        if not filename:
            ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"scrape_{ts}.xlsx"

        filepath = self.output_dir / filename

        # Write all sheets via pandas first
        with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
            for sheet_name, rows in data.items():
                if not rows:
                    continue
                df = pd.DataFrame(rows)
                # Truncate sheet name (Excel limit: 31 chars)
                safe_name = sheet_name[:31]
                df.to_excel(writer, sheet_name=safe_name, index=False)

        # Now apply formatting
        self._format(filepath)
        logger.success(f"XLSX saved: {filepath}")
        return filepath

    def build_from_scrape_results(
        self,
        results: list,
        query:   str = "",
        extra_sheets: dict = None,
    ) -> Path:
        """
        Convenience builder — takes list of ExtractedData objects
        and organises into sheets automatically.
        """
        from core.extractor import ExtractedData

        all_emails   = []
        all_phones   = []
        all_links    = []
        all_social   = []
        all_pages    = []
        all_prices   = []

        for r in results:
            if not isinstance(r, ExtractedData):
                continue

            all_pages.append({
                "title":        r.title,
                "url":          r.url,
                "description":  r.description[:200] if r.description else "",
                "keywords_hit": ", ".join(r.keywords_found),
                "headings":     " | ".join(r.headings[:5]),
            })
            for e in r.emails:
                all_emails.append({"email": e, "source_url": r.url, "source_title": r.title})
            for p in r.phones:
                all_phones.append({"phone": p, "source_url": r.url, "source_title": r.title})
            for l in r.links:
                all_links.append({"url": l, "found_on": r.url})
            for s in r.social_links:
                all_social.append({"social_url": s, "found_on": r.url})
            for pr in r.prices:
                all_prices.append({"price": pr, "source_url": r.url, "source_title": r.title})

        sheets = {
            "Pages":        all_pages,
            "Emails":       all_emails,
            "Phones":       all_phones,
            "Prices":       all_prices,
            "Social Links": all_social,
            "All Links":    all_links,
        }

        if extra_sheets:
            sheets.update(extra_sheets)

        # Remove empty sheets
        sheets = {k: v for k, v in sheets.items() if v}

        ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_q   = "".join(c for c in query if c.isalnum() or c in " _-")[:30].strip()
        filename = f"{safe_q}_{ts}.xlsx" if safe_q else f"scrape_{ts}.xlsx"

        return self.build(sheets, filename)

    def build_jobs_report(self, jobs: list[dict], query: str = "") -> Path:
        ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"jobs_{query[:20]}_{ts}.xlsx"
        return self.build({"Jobs": jobs}, filename)

    def build_leads_report(self, leads: list[dict], industry: str = "") -> Path:
        ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"leads_{industry[:20]}_{ts}.xlsx"
        return self.build({"Leads": leads}, filename)

    def build_market_report(self, research: dict) -> Path:
        ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
        topic    = research.get("topic","market")[:20]
        filename = f"market_{topic}_{ts}.xlsx"
        sheets   = {}
        if research.get("web_results"):
            sheets["Web Results"]   = research["web_results"]
        if research.get("news"):
            sheets["News"]          = research["news"]
        if research.get("reddit"):
            sheets["Reddit"]        = research["reddit"]
        if research.get("prices"):
            sheets["Prices"]        = research["prices"]
        return self.build(sheets, filename)

    # ── Formatting ────────────────────────────────────────────────────────────

    def _format(self, filepath: Path):
        wb = load_workbook(filepath)
        for ws in wb.worksheets:
            self._format_sheet(ws)
        wb.save(filepath)

    def _format_sheet(self, ws):
        if ws.max_row < 1:
            return

        # Header row
        for cell in ws[1]:
            cell.font      = HEADER_FONT
            cell.fill      = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border    = THIN_BORDER

        ws.row_dimensions[1].height = 28

        # Data rows
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            fill = ALT_ROW_FILL if row_idx % 2 == 0 else None
            for cell in row:
                cell.font      = BODY_FONT
                cell.border    = THIN_BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if fill:
                    cell.fill = fill
            ws.row_dimensions[row_idx].height = 20

        # Auto column width (capped at 60)
        for col in ws.columns:
            max_len = max(
                (len(str(cell.value)) for cell in col if cell.value),
                default=10,
            )
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)

        # Freeze header
        ws.freeze_panes = "A2"

        # Auto-filter
        ws.auto_filter.ref = ws.dimensions
